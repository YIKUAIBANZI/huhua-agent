import base64
import copy
import io
import sys
import unittest
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image
from pydantic import ValidationError

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.schemas.editor import ResumeData  # noqa: E402
from app.services.resume_sections import get_resume_sections  # noqa: E402
from app.tools.resume_data_docx import resume_data_to_docx  # noqa: E402


def section(kind, title, content="", section_id=None):
    return {"id": section_id or kind, "kind": kind, "title": title, "content": content}


def paragraphs(data):
    return [p.text for p in Document(io.BytesIO(resume_data_to_docx(data))).paragraphs]


class ResumeSectionsTest(unittest.TestCase):
    def setUp(self):
        self.data = ResumeData.model_validate({
            "basic_info": {"name": "候选人甲", "phone": "13800000000", "email": "a@example.com"},
            "education": [{"school": "甲大学", "major": "计算机", "degree": "本科", "highlights": "教育亮点"}],
            "work_experience": [{"company": "甲公司", "title": "开发实习生", "bullets": ["实习成果"]}],
            "projects": [{"name": "甲项目", "role": "独立开发", "description": "项目原文"}],
            "skills": ["Python", "Java"],
            "certificates": ["英语六级"],
            "self_evaluation": "自我评价原文",
        }).model_dump()

    def test_legacy_defaults_and_project_title_path(self):
        for explicit_none in (False, True):
            data = copy.deepcopy(self.data)
            if not explicit_none:
                data.pop("sections")
            default = get_resume_sections(data)
            self.assertEqual([s["kind"] for s in default], [
                "education", "work_experience", "projects", "skills", "certificates", "self_evaluation",
            ])
            self.assertEqual([s["title"] for s in default], [
                "教育背景", "实习经历", "项目经历", "技能证书", "证书", "自我评价",
            ])
            self.assertEqual(default[2]["title_path"], "/experience_section_title")
            self.assertTrue(all(s["title_path"] is None for i, s in enumerate(default) if i != 2))
            self.assertTrue(all(s["content_path"] is None for s in default))
            text = paragraphs(data)
            self.assertLess(text.index("实习经历"), text.index("项目经历"))
        self.data["experience_section_title"] = "校园经历"
        self.assertEqual(get_resume_sections(self.data)[2]["title"], "校园经历")
        self.assertIn("校园经历", paragraphs(self.data))

    def test_explicit_metadata_keeps_order_titles_and_paths(self):
        self.data["sections"] = [
            section("custom", "补充说明", "第一行\n第二行", "custom-1"),
            section("projects", ""),
            section("education", "学习经历"),
        ]
        resolved = get_resume_sections(self.data)
        self.assertEqual([s["id"] for s in resolved], ["custom-1", "projects", "education"])
        self.assertEqual([s["title"] for s in resolved], ["补充说明", "", "学习经历"])
        self.assertEqual([s["title_path"] for s in resolved], [f"/sections/{i}/title" for i in range(3)])
        self.assertEqual(resolved[0]["content_path"], "/sections/0/content")
        self.assertIsNone(resolved[1]["content_path"])
        text = paragraphs(self.data)
        self.assertLess(text.index("第二行"), text.index("甲项目 / 独立开发"))
        self.assertLess(text.index("项目原文"), text.index("学习经历"))
        self.assertNotIn("项目经历", text)

    def test_all_builtin_kinds_use_their_current_source_fields(self):
        self.data["sections"] = [
            section(kind, f"标题-{kind}") for kind in (
                "self_evaluation", "certificates", "skills", "projects", "work_experience", "education",
            )
        ]
        self.data["skill_groups"] = {"技术栈": ["Rust", "Go"]}
        self.data["projects"][0]["polished_bullets"] = ["最新润色要点"]
        text = paragraphs(self.data)
        for title in [s["title"] for s in self.data["sections"]]:
            self.assertIn(title, text)
        positions = [text.index(s["title"]) for s in self.data["sections"]]
        self.assertEqual(positions, sorted(positions))
        self.assertIn("技术栈：Rust / Go", text)
        self.assertNotIn("Python", text)
        self.assertIn("最新润色要点", text)
        self.assertNotIn("项目原文", text)
        for body in ("自我评价原文", "英语六级", "实习成果", "教育亮点"):
            self.assertIn(body, text)

    def test_custom_multiline_html_is_literal_document_text(self):
        malicious = '<script>alert("x")</script>\r\n<img src=x onerror=alert(1)>\n\nA & B'
        self.data["sections"] = [section("custom", "<b>补充</b>", malicious, "custom-1")]
        blob = resume_data_to_docx(self.data)
        text = [p.text for p in Document(io.BytesIO(blob)).paragraphs]
        self.assertEqual(text[-5:], ["<b>补充</b>", '<script>alert("x")</script>', "<img src=x onerror=alert(1)>", "", "A & B"])
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            xml = archive.read("word/document.xml").decode()
        self.assertIn("&lt;script&gt;", xml)
        self.assertNotIn("<script>", xml)
        self.assertNotIn("<img ", xml)

    def test_custom_title_only_and_body_only(self):
        self.data["sections"] = [
            section("custom", "待填写栏目", "", "custom-1"),
            section("custom", "", "自由正文", "custom-2"),
            section("custom", "", "", "custom-3"),
        ]
        self.assertEqual(paragraphs(self.data)[-2:], ["待填写栏目", "自由正文"])

    def test_removed_builtin_omits_source_content_without_mutating_data(self):
        self.data["sections"] = [section("education", "学习经历")]
        original = copy.deepcopy(self.data)
        text = paragraphs(self.data)
        for omitted in ("项目经历", "实习经历", "甲项目 / 独立开发", "实习成果", "Python", "英语六级", "自我评价原文"):
            self.assertNotIn(omitted, text)
        self.assertIn("学习经历", text)
        self.assertEqual(self.data, original)

    def test_empty_sections_keeps_only_fixed_basic_info_and_photo(self):
        photo = io.BytesIO()
        Image.new("RGB", (4, 4), "white").save(photo, format="PNG")
        self.data["basic_info"]["photo"] = "data:image/png;base64," + base64.b64encode(photo.getvalue()).decode()
        self.data["sections"] = []
        self.assertEqual(get_resume_sections(self.data), [])
        doc = Document(io.BytesIO(resume_data_to_docx(self.data)))
        self.assertEqual([p.text for p in doc.paragraphs if p.text], ["候选人甲", "13800000000 · a@example.com"])
        self.assertEqual(len(doc.inline_shapes), 1)
        self.assertEqual(doc.styles["Normal"].font.name, "微软雅黑")

    def test_empty_builtin_body_is_hidden(self):
        data = ResumeData(sections=[section("education", "空教育")]).model_dump()
        self.assertNotIn("空教育", paragraphs(data))

    def test_legacy_schema_and_metadata_roundtrip(self):
        legacy = copy.deepcopy(self.data)
        legacy.pop("sections")
        self.assertIsNone(ResumeData.model_validate(legacy).sections)
        legacy["sections"] = [section("custom", "", "正文", "custom-1")]
        self.assertEqual(ResumeData.model_validate(legacy).model_dump()["sections"], legacy["sections"])
        legacy["sections"] = []
        self.assertEqual(ResumeData.model_validate(legacy).sections, [])

    def test_duplicate_and_blank_ids_and_duplicate_builtin_kinds_rejected(self):
        invalid = [
            [section("education", "教育", section_id="same"), section("projects", "项目", section_id="same")],
            [section("education", "教育一", section_id="edu-1"), section("education", "教育二", section_id="edu-2")],
            [section("custom", "补充", section_id=" ")],
            [{"id": "", "kind": "custom", "title": "补充"}],
        ]
        for sections in invalid:
            with self.subTest(sections=sections), self.assertRaises(ValidationError):
                ResumeData(sections=sections)
        valid = ResumeData(sections=[section("custom", "一", section_id="one"), section("custom", "二", section_id="two")])
        self.assertEqual(len(valid.sections), 2)

    def test_metadata_limits_and_known_kinds(self):
        invalid = [
            [section("custom", "x" * 101)],
            [section("custom", "标题", "x" * 20001)],
            [section("custom", "标题", section_id="x" * 129)],
            [section("unknown", "标题")],
            [section("custom", "标题", section_id=f"custom-{i}") for i in range(31)],
        ]
        for sections in invalid:
            with self.subTest(first=sections[0]["kind"]), self.assertRaises(ValidationError):
                ResumeData(sections=sections)

    def test_exports_do_not_share_users_section_data(self):
        other = copy.deepcopy(self.data)
        self.data["sections"] = [section("custom", "甲标题", "甲正文", "custom-1")]
        other["basic_info"]["name"] = "候选人乙"
        other["sections"] = [section("projects", "乙项目标题")]
        other["projects"][0]["name"] = "乙项目"
        first, second, repeated = paragraphs(self.data), paragraphs(other), paragraphs(self.data)
        self.assertEqual(first, repeated)
        self.assertIn("甲正文", first)
        self.assertNotIn("甲正文", second)
        self.assertIn("乙项目 / 独立开发", second)
        self.assertNotIn("乙项目 / 独立开发", first)


if __name__ == "__main__":
    unittest.main()

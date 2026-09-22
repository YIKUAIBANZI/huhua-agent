"""
文件文本提取：PDF / DOCX / TXT / MD。
图片类暂不支持（由 chat.py 层拒绝）。
防爆上下文：所有文本统一截断到 MAX_CHARS。
"""

from __future__ import annotations

import io
from pathlib import Path
from zipfile import ZipFile

MAX_CHARS = 6000  # 单个文件提取后文本最大字符数（防爆上下文）
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 20
MAX_DOCX_PARTS = 512
MAX_DOCX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
SUPPORTED_EXT = {".pdf", ".docx", ".txt", ".md", ".markdown"}


class UnsupportedFileError(Exception):
    pass


class FileLimitError(Exception):
    pass


def _truncate(text: str, filename: str) -> str:
    text = text.strip()
    if len(text) <= MAX_CHARS:
        return text
    return (
        text[:MAX_CHARS] + f"\n\n...（{filename} 内容过长，已截断到 {MAX_CHARS} 字符）"
    )


def _extract_pdf(raw: bytes) -> str:
    try:
        import pdfplumber  # type: ignore
    except ImportError as e:
        raise RuntimeError("缺少 pdfplumber 依赖") from e
    chunks: list[str] = []
    char_count = 0
    with pdfplumber.open(io.BytesIO(raw)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise FileLimitError(f"PDF 超过 {MAX_PDF_PAGES} 页处理上限")
        for page in pdf.pages:
            t = page.extract_text() or ""
            if t.strip():
                chunks.append(t)
                char_count += len(t)
                if char_count >= MAX_CHARS:
                    break
    return "\n\n".join(chunks)


def _extract_docx(raw: bytes) -> str:
    from docx import Document

    with ZipFile(io.BytesIO(raw)) as archive:
        parts = archive.infolist()
        if (
            len(parts) > MAX_DOCX_PARTS
            or sum(part.file_size for part in parts) > MAX_DOCX_UNCOMPRESSED_BYTES
        ):
            raise FileLimitError("DOCX 解压后内容过大，请精简文件")
    doc = Document(io.BytesIO(raw))
    parts: list[str] = []
    char_count = 0
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
            char_count += len(t)
            if char_count >= MAX_CHARS:
                return "\n".join(parts)
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join((cell.text or "").strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
                char_count += len(row_text)
                if char_count >= MAX_CHARS:
                    return "\n".join(parts)
    return "\n".join(parts)


def _extract_text(raw: bytes) -> str:
    for encoding in ("utf-8", "gb18030", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(filename: str, raw: bytes) -> str:
    """根据扩展名派发到对应提取器，返回截断后的纯文本。"""
    if len(raw) > MAX_FILE_BYTES:
        raise FileLimitError("文件超过 10MB 上限")
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _extract_pdf(raw)
    elif ext == ".docx":
        text = _extract_docx(raw)
    elif ext in {".txt", ".md", ".markdown"}:
        text = _extract_text(raw)
    else:
        raise UnsupportedFileError(
            f"不支持的文件类型: {ext}。当前支持：PDF / DOCX / TXT / MD"
        )
    return _truncate(text, filename)

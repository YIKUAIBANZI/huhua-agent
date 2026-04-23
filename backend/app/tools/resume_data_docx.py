"""
从 ResumeData dict 生成 Word 文档（不走 markdown 中间层）。

与 docx_exporter.py 的区别：
  - docx_exporter: 输入聊天产出的 markdown 文本（resume_draft）
  - 本模块:        输入结构化 ResumeData（前端表单 / 聊天 state.resume_data）
"""

import base64
import io
import logging
from typing import Any

from docx import Document
from docx.shared import Inches, Pt

logger = logging.getLogger(__name__)


def _decode_photo_data_url(data_url: str) -> bytes | None:
    """解析 data:image/...;base64,xxx 成 bytes。格式错返回 None。"""
    if not data_url or not data_url.startswith("data:image/"):
        return None
    try:
        _, b64 = data_url.split(",", 1)
        return base64.b64decode(b64)
    except Exception as e:
        logger.warning(f"[docx/photo] 解码失败: {e}")
        return None


def _add_section_heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


def _add_kv_line(doc, label: str, value: str) -> None:
    if not value:
        return
    p = doc.add_paragraph()
    run = p.add_run(f"{label}：{value}")
    run.font.size = Pt(10.5)


def resume_data_to_docx(data: dict[str, Any]) -> bytes:
    """把 ResumeData dict 转为 docx bytes。空字段不输出对应行。"""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    basic = data.get("basic_info") or {}

    # 头像（有则嵌入；doc.add_picture 放 inline 段落，右对齐看视觉更像简历照片位）
    photo_bytes = _decode_photo_data_url(basic.get("photo") or "")
    if photo_bytes:
        try:
            photo_p = doc.add_paragraph()
            run = photo_p.add_run()
            run.add_picture(io.BytesIO(photo_bytes), width=Inches(1.2))
            # 右对齐照片段
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            photo_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        except Exception as e:
            logger.warning(f"[docx/photo] add_picture 失败: {e}")

    # 标题：姓名（大号）
    if basic.get("name"):
        title = doc.add_paragraph()
        run = title.add_run(basic["name"])
        run.bold = True
        run.font.size = Pt(18)

    # 基本信息一行拼起来
    meta_parts = []
    for k in ("phone", "email", "location", "objective"):
        if basic.get(k):
            meta_parts.append(basic[k])
    if meta_parts:
        p = doc.add_paragraph()
        p.add_run(" · ".join(meta_parts))

    # 额外 basic_info：身高/年龄/政治面貌/民族（仅显示非空）
    opt_parts = []
    opt_labels = {
        "height": "身高",
        "age": "年龄",
        "gender": "性别",
        "political": "政治面貌",
        "ethnicity": "民族",
        "hometown": "籍贯",
        "birthday": "出生年月",
    }
    for k, label in opt_labels.items():
        if basic.get(k):
            opt_parts.append(f"{label} {basic[k]}")
    if opt_parts:
        doc.add_paragraph(" | ".join(opt_parts))

    # 教育
    if data.get("education"):
        _add_section_heading(doc, "教育背景")
        for edu in data["education"]:
            line = " / ".join(
                x
                for x in [
                    edu.get("school"),
                    edu.get("major"),
                    edu.get("degree"),
                    f"{edu.get('start_date', '')}—{edu.get('end_date', '')}".strip("—"),
                ]
                if x and x != "—"
            )
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            if edu.get("highlights"):
                doc.add_paragraph(edu["highlights"])

    # 经历（projects + work_experience 合并）
    section_title = data.get("experience_section_title") or "经历"
    has_exp = data.get("projects") or data.get("work_experience")
    if has_exp:
        _add_section_heading(doc, section_title)
        # 工作/实习
        for job in data.get("work_experience") or []:
            line = " / ".join(
                x
                for x in [
                    job.get("company"),
                    job.get("title"),
                    f"{job.get('start_date', '')}—{job.get('end_date', '')}".strip("—"),
                ]
                if x and x != "—"
            )
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            for b in job.get("bullets") or []:
                doc.add_paragraph(b, style="List Bullet")
        # 项目
        for proj in data.get("projects") or []:
            line = " / ".join(
                x
                for x in [
                    proj.get("name"),
                    proj.get("role"),
                    f"{proj.get('start_date', '')}—{proj.get('end_date', '')}".strip(
                        "—"
                    ),
                ]
                if x and x != "—"
            )
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            bullets = proj.get("polished_bullets") or []
            if bullets:
                for b in bullets:
                    doc.add_paragraph(b, style="List Bullet")
            elif proj.get("description"):
                doc.add_paragraph(proj["description"])

    # 技能（skill_groups 优先，fallback 到 flat skills）
    skill_groups = data.get("skill_groups") or {}
    if skill_groups and any(v for v in skill_groups.values()):
        _add_section_heading(doc, "技能证书")
        for group_name, items in skill_groups.items():
            if items:
                p = doc.add_paragraph()
                p.add_run(f"{group_name}：").bold = True
                p.add_run(" / ".join(items))
    elif data.get("skills"):
        _add_section_heading(doc, "技能证书")
        for s in data["skills"]:
            doc.add_paragraph(s, style="List Bullet")

    if data.get("certificates"):
        _add_section_heading(doc, "证书")
        for c in data["certificates"]:
            doc.add_paragraph(c, style="List Bullet")

    if data.get("self_evaluation"):
        _add_section_heading(doc, "自我评价")
        doc.add_paragraph(data["self_evaluation"])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
